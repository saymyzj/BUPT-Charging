"""Acceptance console xlsx export regression tests."""

import os
import sys
import tempfile
import unittest

from flask import Flask
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.services.acceptance_service import (
    DEFAULT_START_TIME,
    build_snapshot,
    execute_all,
    execute_event,
    export_table_xlsx,
    initialize_acceptance_database,
    replace_events,
    reset_acceptance,
    save_snapshot,
    snapshot_history,
)
from app.utils.db import execute_db, init_db, query_db


class AcceptanceExportTests(unittest.TestCase):
    def setUp(self):
        self.db_fd, self.db_path = tempfile.mkstemp(suffix=".db")
        self.log_dir = tempfile.mkdtemp()
        self.app = Flask(__name__)
        self.app.config.update(
            TESTING=True,
            DATABASE_PATH=self.db_path,
            LOG_DIR=self.log_dir,
            SECRET_KEY="test-secret",
            JWT_EXPIRATION_HOURS=24,
            AUTO_REBUILD_INCOMPATIBLE_DB=True,
            FAST_CHARGING_PILE_NUM=3,
            TRICKLE_CHARGING_PILE_NUM=2,
            CHARGING_QUEUE_LEN=3,
            WAITING_AREA_SIZE=10,
            DISPATCH_MODE="NORMAL",
            FAULT_DISPATCH_MODE="PRIORITY",
        )
        self.ctx = self.app.app_context()
        self.ctx.push()
        init_db(self.app)

    def tearDown(self):
        self.ctx.pop()
        os.close(self.db_fd)
        os.unlink(self.db_path)

    def test_export_uses_table_view_columns(self):
        reset_acceptance("sample.xlsx")
        replace_events(
            [
                {
                    "event_id": "EVT0001",
                    "source": "XLSX",
                    "at": DEFAULT_START_TIME,
                    "event_type": "APPLY",
                    "vehicle_code": "V1",
                    "charge_mode": "FAST",
                    "value": 20,
                    "raw_text": "(A,V1,F,20)",
                    "enabled": True,
                }
            ],
            "sample.xlsx",
        )
        snapshot = build_snapshot(DEFAULT_START_TIME, phase="AFTER")
        save_snapshot("EVT0001", DEFAULT_START_TIME, "AFTER", snapshot)

        workbook = load_workbook(export_table_xlsx(), data_only=True)
        sheet = workbook.active

        self.assertEqual(sheet.title, "表格视图")
        self.assertEqual(sheet.max_column, 8)
        self.assertEqual(sheet.cell(1, 3).value, "(车号,已充电量,当前费用)")
        self.assertEqual(sheet.cell(1, 8).value, "等候区(车号,充电类型,充电量)；故障队列(车号,已充电量,当前费用)")
        self.assertEqual(
            [sheet.cell(2, column).value for column in range(1, 9)],
            ["时刻", "事件", "快充1", "快充2", "快充3", "慢充1", "慢充2", "等候区"],
        )
        self.assertEqual(sheet.cell(3, 1).value, "06:00:00")
        self.assertEqual(sheet.cell(3, 2).value, "(A,V1,F,20)")

    def test_export_uses_tuple_cells_for_station_queues(self):
        reset_acceptance("sample.xlsx")
        replace_events(
            [
                {
                    "event_id": "EVT0001",
                    "source": "XLSX",
                    "at": DEFAULT_START_TIME,
                    "event_type": "APPLY",
                    "vehicle_code": "V1",
                    "charge_mode": "FAST",
                    "value": 20,
                    "raw_text": "(A,V1,F,20)",
                    "enabled": True,
                }
            ],
            "sample.xlsx",
        )
        execute_event("EVT0001")

        workbook = load_workbook(export_table_xlsx(), data_only=True)
        sheet = workbook.active

        self.assertEqual(sheet.cell(3, 3).value, "(V1,0,0)")
        self.assertNotIn("充电中", sheet.cell(3, 3).value)
        self.assertNotIn("费用", sheet.cell(3, 3).value)

    def test_acceptance_initialization_uses_runtime_station_config(self):
        self.app.config.update(
            FAST_CHARGING_PILE_NUM=2,
            TRICKLE_CHARGING_PILE_NUM=3,
            CHARGING_QUEUE_LEN=4,
            WAITING_AREA_SIZE=9,
            DISPATCH_MODE="EXT_SINGLE_BATCH",
            FAULT_DISPATCH_MODE="PRIORITY",
        )

        initialize_acceptance_database(3)

        station_rows = query_db("SELECT station_code, queue_capacity FROM charging_station ORDER BY station_code")
        self.assertEqual(
            [row["station_code"] for row in station_rows],
            ["FAST_01", "FAST_02", "SLOW_01", "SLOW_02", "SLOW_03"],
        )
        self.assertEqual({int(row["queue_capacity"]) for row in station_rows}, {4})
        config_rows = {
            row["config_key"]: row["config_value"]
            for row in query_db("SELECT config_key, config_value FROM scheduler_config")
        }
        self.assertEqual(config_rows["fast_station_count"], "2")
        self.assertEqual(config_rows["slow_station_count"], "3")
        self.assertEqual(config_rows["waiting_area_capacity"], "9")
        self.assertEqual(config_rows["charging_queue_len"], "4")
        self.assertEqual(config_rows["dispatch_mode"], "EXT_SINGLE_BATCH")
        self.assertEqual(config_rows["fault_dispatch_mode"], "PRIORITY")

        workbook = load_workbook(export_table_xlsx(), data_only=True)
        sheet = workbook.active
        self.assertEqual(
            [sheet.cell(2, column).value for column in range(1, 9)],
            ["时刻", "事件", "快充1", "快充2", "慢充1", "慢充2", "慢充3", "等候区"],
        )

    def test_startup_auto_initialization_rebuilds_runtime_data_from_config(self):
        user_pk = execute_db(
            """
            INSERT INTO user (user_id, username, password_hash, battery_capacity, role)
            VALUES ('U900', 'startup_user', 'hash', 80.0, 'USER')
            """
        )
        execute_db(
            """
            INSERT INTO charge_request (
                request_id, user_id, charge_mode, request_energy, request_status, queue_number, request_time
            ) VALUES ('REQ9000', ?, 'FAST', 20.0, 'WAITING_AREA', 'F9000', '2026-05-20T06:00:00')
            """,
            [user_pk],
        )
        self.app.config.update(
            AUTO_INIT_DB_ON_START=True,
            FAST_CHARGING_PILE_NUM=2,
            TRICKLE_CHARGING_PILE_NUM=1,
            CHARGING_QUEUE_LEN=5,
            WAITING_AREA_SIZE=8,
        )

        init_db(self.app)

        request_count = query_db("SELECT COUNT(*) AS cnt FROM charge_request", one=True)["cnt"]
        station_rows = query_db("SELECT station_code, queue_capacity FROM charging_station ORDER BY station_code")
        self.assertEqual(request_count, 0)
        self.assertEqual([row["station_code"] for row in station_rows], ["FAST_01", "FAST_02", "SLOW_01"])
        self.assertEqual({int(row["queue_capacity"]) for row in station_rows}, {5})
        self.assertEqual(
            query_db("SELECT config_value FROM scheduler_config WHERE config_key = 'waiting_area_capacity'", one=True)[
                "config_value"
            ],
            "8",
        )

    def test_execute_all_preserves_snapshot_history(self):
        initialize_acceptance_database(3)
        replace_events(
            [
                {
                    "event_id": "EVT0001",
                    "source": "XLSX",
                    "at": DEFAULT_START_TIME,
                    "event_type": "APPLY",
                    "vehicle_code": "V1",
                    "charge_mode": "FAST",
                    "value": 20,
                    "raw_text": "(A,V1,F,20)",
                    "enabled": True,
                },
                {
                    "event_id": "EVT0002",
                    "source": "XLSX",
                    "at": "2026-05-20T06:10:00",
                    "event_type": "APPLY",
                    "vehicle_code": "V2",
                    "charge_mode": "SLOW",
                    "value": 10,
                    "raw_text": "(A,V2,T,10)",
                    "enabled": True,
                },
            ],
            "sample.xlsx",
        )

        execute_all()

        history = snapshot_history()
        phases_by_event = {}
        for item in history:
            phases_by_event.setdefault(item["event_id"], set()).add(item["phase"])
        self.assertEqual(phases_by_event["EVT0001"], {"BEFORE", "AFTER"})
        self.assertEqual(phases_by_event["EVT0002"], {"BEFORE", "AFTER"})
        self.assertIn("FINAL", phases_by_event[None])


if __name__ == "__main__":
    unittest.main()
