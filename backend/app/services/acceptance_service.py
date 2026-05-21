"""Acceptance-console services: simulated time, events, execution, and snapshots."""

from __future__ import annotations

import json
import re
from io import BytesIO
from datetime import date, datetime, time, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from flask import current_app

from app.enums import ChargeMode, RequestStatus
from app.services.billing_service import calculate_charge_fee, ensure_request_detail
from app.services.queue_model import (
    fault_dispatch_mode,
    fault_queue_candidates,
    handle_station_fault,
    handle_station_recover,
    log_request_lifecycle,
    refresh_station_after_queue_change,
    run_dispatch_scheduler,
    run_normal_scheduler,
    set_dispatch_mode,
    set_fault_dispatch_mode,
    settle_station_until_time,
)
from app.utils.auth import hash_password
from app.utils.auth import generate_token
from app.utils.db import execute_db, query_db

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # pragma: no cover - surfaced as a user-facing validation error.
    Workbook = None
    load_workbook = None


ACCEPTANCE_BASE_DATE = "2026-05-20"
DEFAULT_START_TIME = f"{ACCEPTANCE_BASE_DATE}T06:00:00"
EVENT_RE = re.compile(r"^\(\s*([ABC])\s*,\s*([^,]+)\s*,\s*([FTO])\s*,\s*([-+]?\d+(?:\.\d+)?)\s*\)$")
TABLE_EXPORT_COLUMNS = [
    ("time", "时刻"),
    ("event", "事件"),
    ("FAST_01", "快充1"),
    ("FAST_02", "快充2"),
    ("FAST_03", "快充3"),
    ("SLOW_01", "慢充1"),
    ("SLOW_02", "慢充2"),
    ("waiting_area", "等候区"),
]


def _parse_dt(value: Any) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, time):
        return datetime.combine(date.fromisoformat(ACCEPTANCE_BASE_DATE), value)
    text = str(value).strip().replace("Z", "+00:00")
    if re.fullmatch(r"\d{2}:\d{2}(:\d{2})?", text):
        fmt = "%H:%M:%S" if text.count(":") == 2 else "%H:%M"
        return datetime.combine(date.fromisoformat(ACCEPTANCE_BASE_DATE), datetime.strptime(text, fmt).time())
    return datetime.fromisoformat(text).replace(tzinfo=None)


def _db_string(value: Any) -> str:
    return _parse_dt(value).strftime("%Y-%m-%dT%H:%M:%S")


def _clock_string(value: Any) -> str:
    return _parse_dt(value).strftime("%H:%M:%S")


def _config_set(key: str, value: Any) -> None:
    execute_db(
        """
        INSERT INTO scheduler_config (config_key, config_value)
        VALUES (?, ?)
        ON CONFLICT(config_key) DO UPDATE SET
            config_value = excluded.config_value,
            updated_at = CURRENT_TIMESTAMP
        """,
        [key, str(value)],
    )


def _config_get(key: str, default: Any = "") -> str:
    row = query_db("SELECT config_value FROM scheduler_config WHERE config_key = ?", [key], one=True)
    return str(row["config_value"]) if row else str(default)


def acceptance_enabled() -> bool:
    return _config_get("acceptance_enabled", "0") == "1"


def acceptance_time(default: Any = None) -> str:
    return _config_get("acceptance_simulation_time", default or DEFAULT_START_TIME)


def _advance_runtime_to(value: Any) -> str:
    parsed = _db_string(value)
    for station in query_db("SELECT id FROM charging_station WHERE station_status = 'RUNNING'"):
        settle_station_until_time(int(station["id"]), parsed)
    run_dispatch_scheduler(parsed)
    return parsed


def set_acceptance_time(value: Any, *, advance_runtime: bool = False) -> str:
    parsed = _db_string(value)
    if advance_runtime and acceptance_enabled():
        _advance_runtime_to(parsed)
    _config_set("acceptance_simulation_time", parsed)
    return parsed


def _next_request_id() -> str:
    row = query_db(
        """
        SELECT request_id
        FROM charge_request
        WHERE request_id LIKE 'REQ%'
        ORDER BY CAST(SUBSTR(request_id, 4) AS INTEGER) DESC
        LIMIT 1
        """,
        one=True,
    )
    next_number = 1 if not row else int(str(row["request_id"])[3:]) + 1
    return f"REQ{next_number:04d}"


def _next_queue_number(charge_mode: str) -> str:
    prefix = "F" if charge_mode == ChargeMode.FAST.value else "T"
    row = query_db(
        """
        SELECT queue_number
        FROM charge_request
        WHERE queue_number LIKE ?
        ORDER BY CAST(SUBSTR(queue_number, 2) AS INTEGER) DESC
        LIMIT 1
        """,
        [f"{prefix}%"],
        one=True,
    )
    next_number = 1 if not row else int(str(row["queue_number"])[1:]) + 1
    return f"{prefix}{next_number}"


def _next_waiting_order(charge_mode: str) -> int:
    row = query_db(
        """
        SELECT COALESCE(MAX(waiting_area_order), 0) AS max_order
        FROM charge_request
        WHERE charge_mode = ?
          AND request_status = ?
        """,
        [charge_mode, RequestStatus.WAITING_AREA.value],
        one=True,
    )
    return int(row["max_order"]) + 1


def _ensure_vehicle_user(vehicle_code: str) -> int:
    vehicle_code = vehicle_code.upper()
    row = query_db("SELECT id FROM user WHERE user_id = ?", [vehicle_code], one=True)
    if row:
        return int(row["id"])
    username = _vehicle_username(vehicle_code)
    return int(
        execute_db(
            """
            INSERT INTO user (user_id, username, password_hash, battery_capacity, role)
            VALUES (?, ?, ?, 100.0, 'USER')
            """,
            [vehicle_code, username, hash_password("123456")],
        )
    )


def _vehicle_username(vehicle_code: str) -> str:
    match = re.fullmatch(r"V(\d+)", vehicle_code.upper())
    if match:
        return f"user{int(match.group(1)):02d}"
    return vehicle_code.lower()


def _vehicle_user_id(vehicle_code: str) -> int | None:
    row = query_db("SELECT id FROM user WHERE user_id = ?", [vehicle_code.upper()], one=True)
    return int(row["id"]) if row else None


def _active_request_for_vehicle(vehicle_code: str):
    user_pk = _vehicle_user_id(vehicle_code)
    if not user_pk:
        return None
    return query_db(
        """
        SELECT
            cr.*,
            cs.power_kw,
            cs.station_code
        FROM charge_request cr
        LEFT JOIN charging_station cs ON cs.id = cr.station_id
        WHERE cr.user_id = ?
          AND cr.request_status IN (?, ?, ?)
        ORDER BY cr.id DESC
        LIMIT 1
        """,
        [
            user_pk,
            RequestStatus.WAITING_AREA.value,
            RequestStatus.QUEUED.value,
            RequestStatus.CHARGING.value,
        ],
        one=True,
    )


def _station_code_from_ref(station_ref: str | None) -> str | None:
    if not station_ref:
        return None
    ref = station_ref.upper()
    if ref.startswith("FAST_") or ref.startswith("SLOW_"):
        return ref
    if ref.startswith("F") and ref[1:].isdigit():
        return f"FAST_{int(ref[1:]):02d}"
    if ref.startswith("T") and ref[1:].isdigit():
        return f"SLOW_{int(ref[1:]):02d}"
    return ref


def _xlsx_charge_mode(mode: str | None) -> str | None:
    if mode == "F":
        return ChargeMode.FAST.value
    if mode == "T":
        return ChargeMode.SLOW.value
    return None


def _reset_runtime_tables() -> None:
    for table in (
        "acceptance_snapshot",
        "acceptance_event",
        "request_detail",
        "charging_session",
        "charge_request",
        "charging_station",
        "notification",
        "scheduler_event_log",
    ):
        execute_db(f"DELETE FROM {table}")
    execute_db("DELETE FROM user WHERE username LIKE 'user%'")


def _clear_all_runtime_data() -> None:
    for table in (
        "acceptance_snapshot",
        "acceptance_event",
        "request_detail",
        "charging_session",
        "charge_request",
        "charging_station",
        "notification",
        "scheduler_event_log",
        "user",
    ):
        execute_db(f"DELETE FROM {table}")


def configure_acceptance_scenario() -> None:
    set_dispatch_mode("NORMAL")
    set_fault_dispatch_mode("TIME_ORDER")
    for index in range(1, 4):
        execute_db(
            """
            INSERT INTO charging_station (station_code, charge_mode, power_kw, station_status, queue_capacity)
            VALUES (?, 'FAST', 30.0, 'RUNNING', 3)
            """,
            [f"FAST_{index:02d}"],
        )
    for index in range(1, 3):
        execute_db(
            """
            INSERT INTO charging_station (station_code, charge_mode, power_kw, station_status, queue_capacity)
            VALUES (?, 'SLOW', 10.0, 'RUNNING', 3)
            """,
            [f"SLOW_{index:02d}"],
        )
    config_values = {
        "fast_station_count": 3,
        "slow_station_count": 2,
        "waiting_area_capacity": 10,
        "charging_queue_len": 3,
        "dispatch_mode": "NORMAL",
        "fault_dispatch_mode": "TIME_ORDER",
    }
    for key, value in config_values.items():
        _config_set(key, value)


def reset_acceptance(sample_name: str = "") -> dict:
    _reset_runtime_tables()
    configure_acceptance_scenario()
    _config_set("acceptance_enabled", "1")
    _config_set("acceptance_status", "PAUSED")
    _config_set("acceptance_sample_name", sample_name)
    set_acceptance_time(DEFAULT_START_TIME)
    snapshot = build_snapshot(DEFAULT_START_TIME, phase="CURRENT")
    save_snapshot(None, DEFAULT_START_TIME, "CURRENT", snapshot)
    return acceptance_state()


def enable_acceptance(sample_name: str = "") -> dict:
    _config_set("acceptance_enabled", "1")
    _config_set("acceptance_status", "PAUSED")
    if sample_name:
        _config_set("acceptance_sample_name", sample_name)
    if not _config_get("acceptance_simulation_time"):
        set_acceptance_time(DEFAULT_START_TIME)
    return acceptance_state()


def disable_acceptance() -> dict:
    _config_set("acceptance_enabled", "0")
    _config_set("acceptance_status", "IDLE")
    return acceptance_state()


def initialize_acceptance_database(user_count: int = 22) -> dict:
    _clear_all_runtime_data()
    configure_acceptance_scenario()
    execute_db(
        """
        INSERT INTO user (user_id, username, password_hash, battery_capacity, role)
        VALUES ('ADMIN', 'admin', ?, 100.0, 'ADMIN')
        """,
        [hash_password("123456")],
    )
    created_users = []
    for index in range(1, max(1, int(user_count)) + 1):
        vehicle_code = f"V{index}"
        username = _vehicle_username(vehicle_code)
        execute_db(
            """
            INSERT INTO user (user_id, username, password_hash, battery_capacity, role)
            VALUES (?, ?, ?, 100.0, 'USER')
            """,
            [vehicle_code, username, hash_password("123456")],
        )
        created_users.append({"vehicle_code": vehicle_code, "username": username, "password": "123456"})
    _config_set("acceptance_enabled", "0")
    _config_set("acceptance_status", "IDLE")
    set_acceptance_time(DEFAULT_START_TIME)
    return {
        "admin": {"username": "admin", "password": "123456"},
        "users": created_users,
        "state": acceptance_state(),
    }


def acceptance_state() -> dict:
    users = _acceptance_accounts()
    return {
        "enabled": acceptance_enabled(),
        "simulation_time": acceptance_time(),
        "status": _config_get("acceptance_status", "IDLE"),
        "sample_name": _config_get("acceptance_sample_name", ""),
        "fault_dispatch_mode": fault_dispatch_mode(),
        "event_counts": _event_counts(),
        "accounts": users,
    }


def _acceptance_accounts() -> dict:
    rows = query_db(
        """
        SELECT id, user_id, username, role
        FROM user
        WHERE username = 'admin'
           OR user_id LIKE 'V%'
        ORDER BY role, CAST(SUBSTR(user_id, 2) AS INTEGER)
        """
    )
    admin = None
    users = []
    for row in rows:
        token = generate_token(row["id"], row["username"], row["role"])
        item = {
            "user_id": row["user_id"],
            "username": row["username"],
            "role": row["role"],
            "password": "123456",
            "token": token,
        }
        if row["role"] == "ADMIN":
            admin = item
        else:
            users.append(item)
    return {"admin": admin, "users": users}


def ensure_sample_users(events: list[dict]) -> list[dict]:
    vehicles = sorted(
        {
            str(item.get("vehicle_code") or "").upper()
            for item in events
            if str(item.get("vehicle_code") or "").upper().startswith("V")
        },
        key=lambda code: int(code[1:]) if code[1:].isdigit() else 9999,
    )
    created = []
    for vehicle_code in vehicles:
        _ensure_vehicle_user(vehicle_code)
        created.append({"vehicle_code": vehicle_code, "username": _vehicle_username(vehicle_code), "password": "123456"})
    return created


def _event_counts() -> dict:
    rows = query_db(
        """
        SELECT status, COUNT(*) AS cnt
        FROM acceptance_event
        GROUP BY status
        """
    )
    counts = {"PENDING": 0, "EXECUTED": 0, "FAILED": 0, "SKIPPED": 0}
    for row in rows:
        counts[row["status"]] = int(row["cnt"])
    counts["TOTAL"] = sum(counts.values())
    return counts


def _raw_event_to_payload(raw: str, row_number: int, event_time: Any) -> dict:
    text = str(raw).replace("，", ",").strip()
    match = EVENT_RE.match(text)
    if not match:
        return {
            "event_id": f"EVT{row_number:04d}",
            "source": "XLSX",
            "at": _db_string(event_time),
            "event_type": "APPLY",
            "vehicle_code": None,
            "station_code": None,
            "charge_mode": None,
            "value": None,
            "raw_text": text,
            "warning": "无法识别事件格式",
            "enabled": False,
        }

    event_kind, target, mode, value_text = match.groups()
    target = target.strip().upper()
    value = float(value_text)
    event_type = "APPLY"
    vehicle_code = target if target.startswith("V") else None
    station_code = None
    charge_mode = _xlsx_charge_mode(mode)

    if event_kind == "A" and mode == "O" and value == 0:
        event_type = "CANCEL_OR_STOP"
    elif event_kind == "C":
        event_type = "CHANGE"
    elif event_kind == "B":
        event_type = "FAULT"
        station_code = _station_code_from_ref(target)
        vehicle_code = None
        charge_mode = None

    return {
        "event_id": f"EVT{row_number:04d}",
        "source": "XLSX",
        "at": _db_string(event_time),
        "event_type": event_type,
        "vehicle_code": vehicle_code,
        "station_code": station_code,
        "charge_mode": charge_mode,
        "value": value,
        "raw_text": text,
        "warning": "",
        "enabled": True,
    }


def parse_xlsx_file(file_storage) -> dict:
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed; cannot parse xlsx acceptance sample")

    filename = Path(file_storage.filename or "acceptance.xlsx").name
    with TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir) / filename
        file_storage.save(temp_path)
        wb = load_workbook(temp_path, data_only=True)
    ws = wb["需填写sheet"] if "需填写sheet" in wb.sheetnames else wb[wb.sheetnames[0]]

    events = []
    for row_number in range(1, ws.max_row + 1):
        time_value = ws.cell(row_number, 1).value
        raw_event = ws.cell(row_number, 2).value
        if not time_value or not raw_event:
            continue
        try:
            parsed_time = _db_string(time_value)
        except Exception:
            continue
        try:
            events.append(_raw_event_to_payload(raw_event, row_number, parsed_time))
        except Exception as exc:  # keep malformed rows editable in the UI.
            events.append(
                {
                    "event_id": f"EVT{row_number:04d}",
                    "source": "XLSX",
                    "at": parsed_time,
                    "event_type": "APPLY",
                    "vehicle_code": None,
                    "station_code": None,
                    "charge_mode": None,
                    "value": None,
                    "raw_text": str(raw_event),
                    "warning": f"识别失败：{exc}",
                    "enabled": False,
                }
            )

    warnings = [
        f"{event['clock'] if 'clock' in event else _clock_string(event['at'])} {event['raw_text']}：{event['warning']}"
        for event in events
        if event.get("warning")
    ]
    return {
        "sample_name": filename,
        "scenario": {
            "fast_station_count": 3,
            "slow_station_count": 2,
            "waiting_area_capacity": 10,
            "charging_queue_len": 3,
            "dispatch_mode": "NORMAL",
            "fault_dispatch_mode": "TIME_ORDER",
            "start_time": "06:00:00",
            "end_time": "11:00:00",
        },
        "events": events,
        "warnings": warnings,
    }


def replace_events(events: list[dict], sample_name: str = "") -> dict:
    execute_db("DELETE FROM acceptance_snapshot")
    execute_db("DELETE FROM acceptance_event")
    if sample_name:
        _config_set("acceptance_sample_name", sample_name)

    for index, item in enumerate(events, start=1):
        if item.get("enabled") is False:
            status = "SKIPPED"
        else:
            status = item.get("status") or "PENDING"
        event_id = item.get("event_id") or f"EVT{index:04d}"
        execute_db(
            """
            INSERT INTO acceptance_event (
                event_id, source, event_time, event_type, vehicle_code,
                station_code, charge_mode, event_value, raw_text, status
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                event_id,
                item.get("source") or "XLSX",
                _db_string(item.get("at") or item.get("event_time") or DEFAULT_START_TIME),
                item.get("event_type") or item.get("type") or "APPLY",
                (item.get("vehicle_code") or "").upper() or None,
                _station_code_from_ref(item.get("station_code")),
                item.get("charge_mode"),
                item.get("value"),
                item.get("raw_text") or "",
                status,
            ],
        )
    accepted_events = list_events()
    created_users = ensure_sample_users(accepted_events)
    return {"events": accepted_events, "created_users": created_users, "state": acceptance_state()}


def list_events() -> list[dict]:
    rows = query_db(
        """
        SELECT *
        FROM acceptance_event
        ORDER BY event_time, id
        """
    )
    return [_event_payload(row) for row in rows]


def _event_payload(row) -> dict:
    return {
        "event_id": row["event_id"],
        "source": row["source"],
        "at": _db_string(row["event_time"]),
        "clock": _clock_string(row["event_time"]),
        "event_type": row["event_type"],
        "vehicle_code": row["vehicle_code"],
        "station_code": row["station_code"],
        "charge_mode": row["charge_mode"],
        "value": row["event_value"],
        "raw_text": row["raw_text"],
        "status": row["status"],
        "error_message": row["error_message"],
        "result": json.loads(row["result_payload"]) if row["result_payload"] else None,
    }


def add_manual_event(payload: dict) -> dict:
    event_id = payload.get("event_id") or f"MANUAL{int(datetime.now().timestamp() * 1000)}"
    status = payload.get("status") or "PENDING"
    item = {
        "event_id": event_id,
        "source": "MANUAL",
        "at": payload.get("at") or acceptance_time(),
        "event_type": payload.get("event_type") or "APPLY",
        "vehicle_code": (payload.get("vehicle_code") or "").upper() or None,
        "station_code": _station_code_from_ref(payload.get("station_code")),
        "charge_mode": payload.get("charge_mode"),
        "value": payload.get("value"),
        "raw_text": payload.get("raw_text") or "手动布置事件",
        "status": status,
        "enabled": True,
    }
    existing = query_db(
        """
        SELECT *
        FROM acceptance_event
        WHERE source = 'MANUAL'
          AND event_time = ?
          AND event_type = ?
          AND COALESCE(vehicle_code, '') = COALESCE(?, '')
          AND COALESCE(station_code, '') = COALESCE(?, '')
          AND COALESCE(charge_mode, '') = COALESCE(?, '')
          AND COALESCE(event_value, -999999) = COALESCE(?, -999999)
          AND status = ?
        ORDER BY id DESC
        LIMIT 1
        """,
        [
            _db_string(item["at"]),
            item["event_type"],
            item["vehicle_code"],
            item["station_code"],
            item["charge_mode"],
            item["value"],
            status,
        ],
        one=True,
    )
    if existing:
        return {"event": _event_payload(existing), "events": list_events(), "duplicate": True}
    execute_db(
        """
        INSERT INTO acceptance_event (
            event_id, source, event_time, event_type, vehicle_code, station_code,
            charge_mode, event_value, raw_text, status, result_payload
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            item["event_id"],
            item["source"],
            _db_string(item["at"]),
            item["event_type"],
            item["vehicle_code"],
            item["station_code"],
            item["charge_mode"],
            item["value"],
            item["raw_text"],
            status,
            json.dumps(payload.get("result") or {}, ensure_ascii=False) if payload.get("result") else None,
        ],
    )
    return {"event": item, "events": list_events()}


def record_manual_event(payload: dict, *, status: str = "EXECUTED", result: dict | None = None) -> dict:
    created = add_manual_event({**payload, "status": status, "result": result or {}})
    if status == "EXECUTED" and not created.get("duplicate"):
        event = created.get("event") or {}
        event_id = event.get("event_id")
        event_time = event.get("at") or payload.get("at") or acceptance_time()
        save_snapshot(event_id, event_time, "AFTER", build_snapshot(event_time, phase="AFTER"))
    return created


def _execute_apply(event, event_dt: str) -> dict:
    charge_mode = event["charge_mode"]
    if charge_mode not in {ChargeMode.FAST.value, ChargeMode.SLOW.value}:
        raise ValueError("申请事件缺少有效充电类型")
    energy = float(event["event_value"])
    if energy <= 0:
        raise ValueError("申请电量必须为正数")
    vehicle_code = event["vehicle_code"]
    if not vehicle_code:
        raise ValueError("申请事件缺少车辆编号")

    user_pk = _ensure_vehicle_user(vehicle_code)
    existing = _active_request_for_vehicle(vehicle_code)
    if existing:
        raise ValueError(f"{vehicle_code} 已有活跃请求 {existing['request_id']}")

    request_id = _next_request_id()
    queue_number = _next_queue_number(charge_mode)
    execute_db(
        """
        INSERT INTO charge_request (
            request_id, user_id, charge_mode, request_energy, request_status,
            queue_number, waiting_area_order, request_time
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            request_id,
            user_pk,
            charge_mode,
            energy,
            RequestStatus.WAITING_AREA.value,
            queue_number,
            _next_waiting_order(charge_mode),
            event_dt,
        ],
    )
    log_request_lifecycle(
        request_id,
        "REQUEST_SUBMITTED",
        event_time=event_dt,
        queue_number=queue_number,
        charge_mode=charge_mode,
        request_energy=energy,
    )
    log_request_lifecycle(
        request_id,
        "WAITING_AREA_ENTERED",
        event_time=event_dt,
        queue_number=queue_number,
        charge_mode=charge_mode,
        request_energy=energy,
        description="公共等候区等待调度",
    )
    dispatch_result = run_dispatch_scheduler(event_dt, charge_mode if charge_mode else None)
    return {"request_id": request_id, "queue_number": queue_number, "dispatch": dispatch_result}


def _execute_change(event, event_dt: str) -> dict:
    req = _active_request_for_vehicle(event["vehicle_code"] or "")
    if not req:
        raise ValueError("未找到可修改的活跃请求")

    charge_mode = event["charge_mode"] or req["charge_mode"]
    energy = (
        float(event["event_value"])
        if event["event_value"] is not None and float(event["event_value"]) > 0
        else float(req["request_energy"])
    )
    if req["request_status"] in {RequestStatus.QUEUED.value, RequestStatus.CHARGING.value}:
        if charge_mode != req["charge_mode"]:
            raise ValueError("充电区请求不允许变更充电模式")
        estimated_finish = req["estimated_finish_time"]
        if req["request_status"] == RequestStatus.CHARGING.value and req["charge_start_time"] and req["power_kw"]:
            finish_dt = _parse_dt(req["charge_start_time"]) + timedelta(seconds=int(energy * 3600 / float(req["power_kw"])))
            estimated_finish = _db_string(finish_dt)
        execute_db(
            """
            UPDATE charge_request
            SET request_energy = ?,
                estimated_finish_time = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            [energy, estimated_finish, req["id"]],
        )
        if req["station_id"]:
            settle_station_until_time(int(req["station_id"]), event_dt)
            refresh_station_after_queue_change(int(req["station_id"]), event_dt)
        dispatch_result = run_dispatch_scheduler(event_dt)
        return {"request_id": req["request_id"], "queue_number": req["queue_number"], "dispatch": dispatch_result}

    if req["request_status"] != RequestStatus.WAITING_AREA.value:
        raise ValueError("当前状态不允许修改")

    queue_number = req["queue_number"]
    waiting_order = req["waiting_area_order"]
    if charge_mode != req["charge_mode"]:
        queue_number = _next_queue_number(charge_mode)
        waiting_order = _next_waiting_order(charge_mode)

    execute_db(
        """
        UPDATE charge_request
        SET charge_mode = ?,
            request_energy = ?,
            queue_number = ?,
            waiting_area_order = ?,
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        [charge_mode, energy, queue_number, waiting_order, req["id"]],
    )
    dispatch_result = run_dispatch_scheduler(event_dt)
    return {"request_id": req["request_id"], "queue_number": queue_number, "dispatch": dispatch_result}


def _execute_cancel_or_stop(event, event_dt: str) -> dict:
    req = _active_request_for_vehicle(event["vehicle_code"] or "")
    if not req:
        raise ValueError("未找到可取消或提前结束的活跃请求")

    if req["request_status"] == RequestStatus.WAITING_AREA.value:
        execute_db(
            """
            UPDATE charge_request
            SET request_status = ?,
                waiting_area_order = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            [RequestStatus.CANCELLED.value, req["id"]],
        )
        dispatch_result = run_dispatch_scheduler(event_dt)
        return {"request_id": req["request_id"], "request_status": RequestStatus.CANCELLED.value, "dispatch": dispatch_result}

    if req["station_id"]:
        settle_station_until_time(int(req["station_id"]), event_dt)
        req = _active_request_for_vehicle(event["vehicle_code"] or "")
        if not req:
            raise ValueError("请求已在推进时间时完成")

    stop_dt = _parse_dt(event_dt)
    start_dt = _parse_dt(req["charge_start_time"] or event_dt)
    if req["request_status"] == RequestStatus.CHARGING.value:
        if stop_dt < start_dt:
            stop_dt = start_dt
        duration_seconds = max(0, int((stop_dt - start_dt).total_seconds()))
        actual_energy = min(float(req["request_energy"]), round(float(req["power_kw"]) * duration_seconds / 3600.0, 2))
        start_value = _db_string(start_dt)
    elif req["request_status"] == RequestStatus.QUEUED.value:
        duration_seconds = 0
        actual_energy = 0.0
        start_value = event_dt
    else:
        raise ValueError("当前状态不允许取消或提前结束")

    execute_db(
        """
        UPDATE charge_request
        SET request_status = ?,
            actual_energy = ?,
            charge_start_time = ?,
            charge_stop_time = ?,
            charge_duration_seconds = ?,
            station_queue_position = NULL,
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        [
            RequestStatus.COMPLETED_EARLY.value,
            actual_energy,
            start_value,
            event_dt,
            duration_seconds,
            req["id"],
        ],
    )
    if req["station_id"]:
        execute_db(
            """
            UPDATE charging_session
            SET end_time = ?,
                actual_energy = ?,
                status = ?
            WHERE request_id = ?
            """,
            [event_dt, actual_energy, RequestStatus.COMPLETED_EARLY.value, req["id"]],
        )
        if actual_energy > 0 or duration_seconds > 0:
            execute_db(
                """
                UPDATE charging_station
                SET total_charge_count = total_charge_count + 1,
                    total_charge_seconds = total_charge_seconds + ?,
                    total_charge_energy = total_charge_energy + ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                [duration_seconds, actual_energy, req["station_id"]],
            )
        refresh_station_after_queue_change(int(req["station_id"]), event_dt)
    run_dispatch_scheduler(event_dt)
    ensure_request_detail(req["request_id"])
    return {"request_id": req["request_id"], "request_status": RequestStatus.COMPLETED_EARLY.value}


def _execute_fault(event, event_dt: str) -> dict:
    station_code = event["station_code"]
    if not station_code:
        raise ValueError("故障事件缺少充电桩编号")
    result = handle_station_fault(station_code, event_dt)
    minutes = float(event["event_value"] or 0)
    if minutes > 0:
        recover_time = _db_string(_parse_dt(event_dt) + timedelta(minutes=minutes))
        event_id = f"{event['event_id']}-RECOVER"
        exists = query_db("SELECT id FROM acceptance_event WHERE event_id = ?", [event_id], one=True)
        if not exists:
            execute_db(
                """
                INSERT INTO acceptance_event (
                    event_id, source, event_time, event_type, station_code, raw_text, status
                )
                VALUES (?, 'SYSTEM', ?, 'RECOVER', ?, ?, 'PENDING')
                """,
                [event_id, recover_time, station_code, f"{station_code} 故障恢复"],
            )
    return result or {}


def _execute_recover(event, event_dt: str) -> dict:
    station_code = event["station_code"]
    if not station_code:
        raise ValueError("恢复事件缺少充电桩编号")
    return handle_station_recover(station_code, event_dt) or {}


def execute_event(event_id: str) -> dict:
    row = query_db("SELECT * FROM acceptance_event WHERE event_id = ?", [event_id], one=True)
    if not row:
        raise ValueError("验收事件不存在")
    if row["status"] != "PENDING":
        return {"event": _event_payload(row), "skipped": True}

    event_dt = _db_string(row["event_time"])
    set_acceptance_time(event_dt, advance_runtime=True)
    before = build_snapshot(event_dt, phase="BEFORE")
    save_snapshot(row["event_id"], event_dt, "BEFORE", before)
    try:
        if row["event_type"] == "APPLY":
            result = _execute_apply(row, event_dt)
        elif row["event_type"] == "CHANGE":
            result = _execute_change(row, event_dt)
        elif row["event_type"] == "CANCEL_OR_STOP":
            result = _execute_cancel_or_stop(row, event_dt)
        elif row["event_type"] == "FAULT":
            result = _execute_fault(row, event_dt)
        elif row["event_type"] == "RECOVER":
            result = _execute_recover(row, event_dt)
        else:
            raise ValueError(f"未知事件类型 {row['event_type']}")

        after = build_snapshot(event_dt, phase="AFTER")
        save_snapshot(row["event_id"], event_dt, "AFTER", after)
        execute_db(
            """
            UPDATE acceptance_event
            SET status = 'EXECUTED',
                result_payload = ?,
                error_message = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            [json.dumps(result, ensure_ascii=False), row["id"]],
        )
    except Exception as exc:
        after = build_snapshot(event_dt, phase="AFTER")
        save_snapshot(row["event_id"], event_dt, "AFTER", after)
        execute_db(
            """
            UPDATE acceptance_event
            SET status = 'FAILED',
                error_message = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = ?
            """,
            [str(exc), row["id"]],
        )
        raise

    updated = query_db("SELECT * FROM acceptance_event WHERE id = ?", [row["id"]], one=True)
    return {"event": _event_payload(updated), "before": before, "after": after}


def execute_current() -> dict:
    current = acceptance_time()
    rows = query_db(
        """
        SELECT event_id
        FROM acceptance_event
        WHERE event_time = ?
          AND status = 'PENDING'
        ORDER BY id
        """,
        [current],
    )
    results = []
    for row in rows:
        results.append(execute_event(row["event_id"]))
    _config_set("acceptance_status", "PAUSED")
    return {"simulation_time": current, "results": results, "snapshot": build_snapshot(current)}


def execute_until(target_time: Any) -> dict:
    target = _db_string(target_time)
    results = []
    _config_set("acceptance_status", "RUNNING")
    while True:
        rows = query_db(
            """
            SELECT event_id
            FROM acceptance_event
            WHERE event_time <= ?
              AND status = 'PENDING'
            ORDER BY event_time, id
            """,
            [target],
        )
        if not rows:
            break
        for row in rows:
            try:
                results.append(execute_event(row["event_id"]))
            except Exception as exc:
                results.append({"event_id": row["event_id"], "error": str(exc)})
    set_acceptance_time(target, advance_runtime=True)
    for station in query_db("SELECT id FROM charging_station WHERE station_status = 'RUNNING'"):
        settle_station_until_time(int(station["id"]), target)
    _config_set("acceptance_status", "PAUSED")
    snapshot = build_snapshot(target)
    save_snapshot(None, target, "CURRENT", snapshot)
    return {"simulation_time": target, "results": results, "snapshot": snapshot, "events": list_events()}


def execute_all() -> dict:
    results = []
    result = {"events": list_events(), "snapshot": build_snapshot(acceptance_time())}
    for _ in range(1000):
        last = query_db(
            "SELECT MAX(event_time) AS max_time FROM acceptance_event WHERE status = 'PENDING'",
            one=True,
        )
        target = last["max_time"] if last and last["max_time"] else None
        if not target:
            break
        result = execute_until(target)
        results.extend(result.get("results") or [])

    for _ in range(1000):
        next_time = _next_active_time()
        if not next_time:
            break
        active = _active_count()
        if active <= 0:
            break
        for station in query_db("SELECT id FROM charging_station WHERE station_status = 'RUNNING'"):
            settle_station_until_time(int(station["id"]), next_time)
        set_acceptance_time(next_time, advance_runtime=True)
    final_time = acceptance_time()
    snapshot = build_snapshot(final_time, phase="FINAL")
    save_snapshot(None, final_time, "FINAL", snapshot)
    _config_set("acceptance_status", "COMPLETED")
    return {**result, "results": results, "simulation_time": final_time, "final_snapshot": snapshot, "events": list_events()}


def _active_count() -> int:
    row = query_db(
        """
        SELECT COUNT(*) AS cnt
        FROM charge_request
        WHERE request_status IN (?, ?, ?)
        """,
        [RequestStatus.WAITING_AREA.value, RequestStatus.QUEUED.value, RequestStatus.CHARGING.value],
        one=True,
    )
    return int(row["cnt"]) if row else 0


def _next_active_time() -> str | None:
    row = query_db(
        """
        SELECT MIN(estimated_finish_time) AS next_time
        FROM charge_request
        WHERE request_status = ?
          AND estimated_finish_time IS NOT NULL
        """,
        [RequestStatus.CHARGING.value],
        one=True,
    )
    return row["next_time"] if row and row["next_time"] else None


def _live_energy_and_fee(row, at_time: str) -> tuple[float, float]:
    if row["request_status"] in {RequestStatus.COMPLETED.value, RequestStatus.COMPLETED_EARLY.value, RequestStatus.FAULT_INTERRUPTED.value}:
        energy = float(row["actual_energy"] or 0)
        detail = query_db(
            """
            SELECT total_fee
            FROM request_detail
            WHERE request_id = ?
            ORDER BY id DESC
            LIMIT 1
            """,
            [row["id"]],
            one=True,
        )
        return energy, float(detail["total_fee"]) if detail else 0.0

    if row["request_status"] != RequestStatus.CHARGING.value or not row["charge_start_time"] or not row["power_kw"]:
        return 0.0, 0.0

    start_dt = _parse_dt(row["charge_start_time"])
    at_dt = max(start_dt, _parse_dt(at_time))
    duration = max(0, int((at_dt - start_dt).total_seconds()))
    energy = min(float(row["request_energy"]), round(float(row["power_kw"]) * duration / 3600.0, 2))
    charge_fee = calculate_charge_fee(row["charge_start_time"], at_time, float(row["power_kw"]))
    service_fee = round(energy * 0.8, 2)
    return energy, round(charge_fee + service_fee, 2)


def build_snapshot(snapshot_time: Any | None = None, phase: str = "CURRENT") -> dict:
    at_time = _db_string(snapshot_time or acceptance_time())

    station_rows = query_db(
        """
        SELECT *
        FROM charging_station
        ORDER BY station_code
        """
    )
    stations = []
    for station in station_rows:
        queue_rows = query_db(
            """
            SELECT
                cr.*,
                u.user_id AS vehicle_code,
                cs.power_kw,
                cs.station_code
            FROM charge_request cr
            JOIN user u ON u.id = cr.user_id
            JOIN charging_station cs ON cs.id = cr.station_id
            WHERE cr.station_id = ?
              AND cr.request_status IN (?, ?)
            ORDER BY cr.station_queue_position, cr.id
            """,
            [station["id"], RequestStatus.CHARGING.value, RequestStatus.QUEUED.value],
        )
        queue = []
        for row in queue_rows:
            energy, fee = _live_energy_and_fee(row, at_time)
            queue.append(
                {
                    "vehicle_code": row["vehicle_code"],
                    "request_id": row["request_id"],
                    "queue_number": row["queue_number"],
                    "status": row["request_status"],
                    "position": row["station_queue_position"],
                    "request_energy": float(row["request_energy"]),
                    "charged_energy": energy,
                    "current_fee": fee,
                    "estimated_finish_time": row["estimated_finish_time"],
                }
            )
        stations.append(
            {
                "station_code": station["station_code"],
                "display_name": station["station_code"].replace("FAST_", "快充").replace("SLOW_", "慢充"),
                "charge_mode": station["charge_mode"],
                "station_status": station["station_status"],
                "queue_capacity": int(station["queue_capacity"]),
                "queue_length": len(queue),
                "queue": queue,
                "current": queue[0] if queue and queue[0]["status"] == RequestStatus.CHARGING.value else None,
            }
        )

    waiting_rows = query_db(
        """
        SELECT
            cr.request_id,
            cr.queue_number,
            cr.charge_mode,
            cr.request_energy,
            cr.waiting_area_order,
            u.user_id AS vehicle_code
        FROM charge_request cr
        JOIN user u ON u.id = cr.user_id
        WHERE cr.request_status = ?
          AND cr.waiting_area_order > 0
        ORDER BY cr.waiting_area_order, cr.id
        """,
        [RequestStatus.WAITING_AREA.value],
    )
    waiting_area = [
        {
            "vehicle_code": row["vehicle_code"],
            "request_id": row["request_id"],
            "queue_number": row["queue_number"],
            "charge_mode": row["charge_mode"],
            "request_energy": float(row["request_energy"]),
            "waiting_order": row["waiting_area_order"],
            "is_fault_queue": False,
        }
        for row in waiting_rows
    ]
    fault_queue = [
        {
            "vehicle_code": row["user_id"],
            "request_id": row["request_id"],
            "queue_number": row["queue_number"],
            "source_queue_number": row["source_queue_number"],
            "effective_queue_number": row["effective_queue_number"],
            "charge_mode": row["charge_mode"],
            "request_energy": float(row["request_energy"]),
            "charged_energy": 0.0,
            "current_fee": 0.0,
            "status": "FAULT_QUEUE",
            "waiting_order": row["waiting_area_order"],
            "is_fault_queue": True,
        }
        for row in fault_queue_candidates()
    ]

    pending_at_time = query_db(
        """
        SELECT *
        FROM acceptance_event
        WHERE event_time = ?
        ORDER BY id
        """,
        [at_time],
    )
    integrity = _integrity_report(stations, waiting_area, fault_queue)
    return {
        "snapshot_time": at_time,
        "clock": _clock_string(at_time),
        "phase": phase,
        "acceptance": acceptance_state(),
        "stations": stations,
        "waiting_area": waiting_area,
        "waiting_area_count": len(waiting_area),
        "fault_queue": fault_queue,
        "fault_queue_count": len(fault_queue),
        "events_at_time": [_event_payload(row) for row in pending_at_time],
        "integrity": integrity,
        "table_row": _table_row(at_time, stations, waiting_area, pending_at_time, fault_queue),
    }


def _integrity_report(stations: list[dict], waiting_area: list[dict], fault_queue: list[dict] | None = None) -> dict:
    warnings: list[str] = []
    active_vehicle_codes: list[str] = []
    for station in stations:
        if int(station["queue_length"]) > int(station["queue_capacity"]):
            warnings.append(f"{station['station_code']} 队列超过容量")
        current_count = sum(1 for item in station["queue"] if item["status"] == RequestStatus.CHARGING.value)
        if current_count > 1:
            warnings.append(f"{station['station_code']} 同时存在多个充电中请求")
        active_vehicle_codes.extend(item["vehicle_code"] for item in station["queue"])
    active_vehicle_codes.extend(item["vehicle_code"] for item in waiting_area)
    active_vehicle_codes.extend(item["vehicle_code"] for item in (fault_queue or []))
    duplicated = sorted({code for code in active_vehicle_codes if active_vehicle_codes.count(code) > 1})
    if duplicated:
        warnings.append("车辆重复出现在活跃队列：" + "、".join(duplicated))
    capacity = int(_config_get("waiting_area_capacity", 10) or 10)
    if len(waiting_area) > capacity:
        warnings.append(f"等候区超过容量 {capacity}")
    return {"ok": not warnings, "warnings": warnings}


def _table_row(at_time: str, stations: list[dict], waiting_area: list[dict], events, fault_queue: list[dict] | None = None) -> dict:
    station_map = {station["station_code"]: station for station in stations}
    event_text = "；".join(row["raw_text"] or row["event_id"] for row in events)

    def cell(station_code: str) -> str:
        station = station_map.get(station_code)
        if not station:
            return ""
        if station["station_status"] == "FAULT":
            return "故障"
        if not station["queue"]:
            return "空闲"
        parts = []
        for item in station["queue"]:
            status = "充电中" if item["status"] == RequestStatus.CHARGING.value else "排队"
            position = item.get("position") or "-"
            parts.append(
                f"{item['vehicle_code']} {status} "
                f"已充{item['charged_energy']}度 "
                f"费用¥{item['current_fee']} "
                f"队列{position}"
            )
        return "\n".join(parts)

    waiting = "，".join(
        f"{item['vehicle_code']}({item['charge_mode']}，{item['request_energy']}度，顺序{item['waiting_order']})"
        for item in waiting_area
    )
    fault_waiting = "\n".join(
        f"{item['vehicle_code']} 故障队列 "
        f"已充{_format_export_number(item.get('charged_energy', 0))}度 "
        f"费用¥{_format_export_number(item.get('current_fee', 0))} "
        f"队列{item['effective_queue_number'] or item['queue_number']}"
        for item in (fault_queue or [])
    )
    waiting_text = "\n".join(part for part in [waiting, f"故障队列：\n{fault_waiting}" if fault_waiting else ""] if part)
    return {
        "time": _clock_string(at_time),
        "event": event_text,
        "FAST_01": cell("FAST_01"),
        "FAST_02": cell("FAST_02"),
        "FAST_03": cell("FAST_03"),
        "SLOW_01": cell("SLOW_01"),
        "SLOW_02": cell("SLOW_02"),
        "waiting_area": waiting_text,
    }


def _format_export_number(value: Any) -> str:
    number = float(value or 0)
    rounded = round(number, 2)
    if rounded == int(rounded):
        return str(int(rounded))
    return f"{rounded:.2f}".rstrip("0").rstrip(".")


def _table_row_for_xlsx(at_time: str, stations: list[dict], waiting_area: list[dict], events, fault_queue: list[dict] | None = None) -> dict:
    station_map = {station["station_code"]: station for station in stations}
    event_text = "；".join(row["raw_text"] or row["event_id"] for row in events)

    def cell(station_code: str) -> str:
        station = station_map.get(station_code)
        if not station:
            return ""
        if station["station_status"] == "FAULT":
            return "故障"
        if not station["queue"]:
            return ""
        return "\n".join(
            f"({item['vehicle_code']},{_format_export_number(item['charged_energy'])},{_format_export_number(item['current_fee'])})"
            for item in station["queue"]
        )

    def waiting_mode(charge_mode: str) -> str:
        return "F" if charge_mode == ChargeMode.FAST.value else "T"

    waiting = "\n".join(
        f"({item['vehicle_code']},{waiting_mode(item['charge_mode'])},{_format_export_number(item['request_energy'])})"
        for item in waiting_area
    )
    fault_waiting = "\n".join(
        f"故障队列({item['vehicle_code']},{_format_export_number(item.get('charged_energy', 0))},{_format_export_number(item.get('current_fee', 0))})"
        for item in (fault_queue or [])
    )
    waiting_text = "\n".join(part for part in [waiting, fault_waiting] if part)
    return {
        "time": _clock_string(at_time),
        "event": event_text,
        "FAST_01": cell("FAST_01"),
        "FAST_02": cell("FAST_02"),
        "FAST_03": cell("FAST_03"),
        "SLOW_01": cell("SLOW_01"),
        "SLOW_02": cell("SLOW_02"),
        "waiting_area": waiting_text,
    }


def save_snapshot(event_id: str | None, snapshot_time: Any, phase: str, snapshot: dict) -> None:
    execute_db(
        """
        INSERT INTO acceptance_snapshot (event_id, snapshot_time, snapshot_phase, snapshot_payload)
        VALUES (?, ?, ?, ?)
        """,
        [event_id, _db_string(snapshot_time), phase, json.dumps(snapshot, ensure_ascii=False)],
    )


def snapshot_history() -> list[dict]:
    rows = query_db(
        """
        SELECT id, event_id, snapshot_time, snapshot_phase, snapshot_payload, created_at
        FROM acceptance_snapshot
        ORDER BY id DESC
        LIMIT 200
        """
    )
    history = []
    for row in rows:
        payload = json.loads(row["snapshot_payload"])
        history.append(
            {
                "id": row["id"],
                "event_id": row["event_id"],
                "snapshot_time": _db_string(row["snapshot_time"]),
                "clock": _clock_string(row["snapshot_time"]),
                "phase": row["snapshot_phase"],
                "created_at": row["created_at"],
                "summary": {
                    "station_count": len(payload.get("stations", [])),
                    "waiting_area_count": payload.get("waiting_area_count", 0),
                    "events": len(payload.get("events_at_time", [])),
                },
                "snapshot": payload,
            }
        )
    return history


def _table_row_from_snapshot_payload(payload: dict, fallback_time: Any, *, xlsx_format: bool = False) -> dict | None:
    stations = payload.get("stations")
    waiting_area = payload.get("waiting_area")
    fault_queue = payload.get("fault_queue") or []
    if isinstance(stations, list) and isinstance(waiting_area, list):
        row_builder = _table_row_for_xlsx if xlsx_format else _table_row
        return row_builder(
            payload.get("snapshot_time") or fallback_time,
            stations,
            waiting_area,
            payload.get("events_at_time") or [],
            fault_queue,
        )
    return payload.get("table_row")


def _table_export_rows() -> list[dict]:
    rows = []
    seen_keys = set()
    snapshots = query_db(
        """
        SELECT id, event_id, snapshot_time, snapshot_phase, snapshot_payload
        FROM acceptance_snapshot
        WHERE snapshot_phase = 'AFTER'
        ORDER BY snapshot_time, id
        """
    )
    for row in snapshots:
        payload = json.loads(row["snapshot_payload"])
        table_row = _table_row_from_snapshot_payload(payload, row["snapshot_time"], xlsx_format=True)
        if not table_row:
            continue
        key = row["event_id"] or f"{_db_string(row['snapshot_time'])}-{row['id']}"
        if key in seen_keys:
            continue
        seen_keys.add(key)
        rows.append(table_row)

    current = build_snapshot(acceptance_time())
    current_row = _table_row_from_snapshot_payload(current, acceptance_time(), xlsx_format=True)
    if current_row:
        current_key = (current_row.get("time"), current_row.get("event") or "状态快照")
        if not any((row.get("time"), row.get("event") or "状态快照") == current_key for row in rows):
            rows.append(current_row)
    return rows


def export_table_xlsx() -> BytesIO:
    if Workbook is None:
        raise RuntimeError("openpyxl is not installed; cannot export xlsx")

    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "表格视图"
    ws.append(["", "", "(车号,已充电量,当前费用)", "", "", "", "", "等候区(车号,充电类型,充电量)；故障队列(车号,已充电量,当前费用)"])
    ws.append([label for _, label in TABLE_EXPORT_COLUMNS])

    for row in _table_export_rows():
        ws.append(
            [
                row.get(key) or ("状态快照" if key == "event" else "")
                for key, _ in TABLE_EXPORT_COLUMNS
            ]
        )

    header_fill = PatternFill("solid", fgColor="E8F3EC")
    header_font = Font(bold=True, color="1F3327")
    for header_row in ws.iter_rows(min_row=1, max_row=2):
        for cell in header_row:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = {
        "A": 12,
        "B": 28,
        "C": 30,
        "D": 30,
        "E": 30,
        "F": 30,
        "G": 30,
        "H": 42,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:H{ws.max_row}"

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def _event_export_text(event: dict) -> str:
    if event.get("event_type") == "FAULT":
        station = event.get("station_code") or ""
        index = station.split("_")[-1].lstrip("0") if station else ""
        return f"(B,F{index},{event.get('value') or 0:g})"
    if event.get("event_type") == "RECOVER":
        station = event.get("station_code") or ""
        index = station.split("_")[-1].lstrip("0") if station else ""
        return f"(B,F{index},0)"
    vehicle = event.get("vehicle_code") or ""
    mode = "F" if event.get("charge_mode") == ChargeMode.FAST.value else "T"
    action = "C" if event.get("event_type") == "CHANGE" else "A"
    if event.get("event_type") == "CANCEL_OR_STOP":
        return f"(A,{vehicle},O,0)"
    return f"({action},{vehicle},{mode},{event.get('value') or 0:g})"


def set_status(status: str) -> dict:
    if status not in {"IDLE", "PAUSED", "RUNNING", "COMPLETED"}:
        raise ValueError("invalid acceptance status")
    _config_set("acceptance_status", status)
    return acceptance_state()
